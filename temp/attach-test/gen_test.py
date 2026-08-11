"""attachment_reader 直测：生成 md/docx/xlsx/pdf/png 样例 → build_attachment_context 组装。
运行：python temp/attach-test/gen_test.py
"""
import json
import sys
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
ATT = REPO / "temp" / "attach-test" / "conversations" / "test-conv" / "attachments"
ATT.mkdir(parents=True, exist_ok=True)


def build_minimal_pdf(text: str, path: Path) -> None:
    """手写最小合法 PDF（Helvetica latin-1 文本），含正确 xref 偏移。"""
    content = f"BT /F1 14 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref_pos = len(out)
    out += b"xref\n0 " + str(len(objs) + 1).encode() + b"\n"
    out += b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%EOF\n".encode()
    path.write_bytes(bytes(out))


# --- 生成各类型样例 ---
(ATT / "a1.md").write_text("# 会议纪要\n\n预算 10 万，下周五复盘。", encoding="utf-8")

from docx import Document  # noqa: E402
doc = Document()
doc.add_paragraph("需求文档")
doc.add_paragraph("核心功能：附件发送。")
doc.save(ATT / "a2.docx")

import openpyxl  # noqa: E402
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "数据"
ws.append(["月份", "销售额"])
ws.append(["一月", 100])
ws.append(["二月", 200])
wb.save(ATT / "a3.xlsx")

build_minimal_pdf("Hello NA attachment pdf", ATT / "a4.pdf")

(ATT / "a5.png").write_bytes(b"\x89PNG\r\n\x1a\n" + b"0" * 64)

meta = {
    "conversation_id": "test-conv",
    "attachments": [
        {"id": "m1", "name": "meeting.md", "size": 10, "mime": "text/markdown", "stored_name": "a1.md", "uploaded_at": 1},
        {"id": "d1", "name": "需求.docx", "size": 10, "mime": "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "stored_name": "a2.docx", "uploaded_at": 1},
        {"id": "x1", "name": "sales.xlsx", "size": 10, "mime": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "stored_name": "a3.xlsx", "uploaded_at": 1},
        {"id": "p1", "name": "report.pdf", "size": 10, "mime": "application/pdf", "stored_name": "a4.pdf", "uploaded_at": 1},
        {"id": "i1", "name": "photo.png", "size": 64, "mime": "image/png", "stored_name": "a5.png", "uploaded_at": 1},
    ],
}
(ATT / "attachments.json").write_text(json.dumps(meta, ensure_ascii=False), encoding="utf-8")

# --- 组装并断言 ---
sys.path.insert(0, str(REPO))
from sidecar.services.context.attachment_reader import build_attachment_context  # noqa: E402

ctx = build_attachment_context(ATT.parents[1], "test-conv", ["m1", "d1", "x1", "p1", "i1"])
print(ctx)
print("\n=== 断言 ===")
assert "会议纪要" in ctx, "md 内容缺失"
assert "附件发送" in ctx, "docx 内容缺失"
assert "销售额" in ctx and "一月" in ctx, "xlsx 内容缺失"
assert "Hello NA attachment pdf" in ctx, "pdf 内容缺失"
assert "photo.png" in ctx and "仅引用" in ctx, "png 应为仅引用"
assert "无法解析内容" in ctx, "png 引用说明缺失"
print("全部通过 ✅")
