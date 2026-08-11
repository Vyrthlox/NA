"""附件内容读取与上下文组装。

按扩展名把已上传附件转成可注入 LLM 的输入：
- txt / md：直接解码文本；html：解码后剥离标签、折叠空白
- pdf / docx / xlsx：pypdf / python-docx / openpyxl 提取文本（解析库懒加载，
  未装或缺依赖时降级为仅引用，绝不因单个附件中断本轮对话）
- png / jpg / jpeg：视觉模型场景下读取为 base64 data URL，随本轮用户消息以
  OpenAI 兼容 image_url content-part 传入；超出体积/数量上限或读取失败时降级
  为仅引用（文件名 + 类型 + 大小），绝不因图片问题中断本轮对话

调用方（/chat 路由）在 prepare_turn 前调用 build_attachment_context 组装文本块、
build_image_parts 组装图片 parts；空附件 / 附件不可读时返回空值，不注入任何上下文。
"""

from __future__ import annotations

import base64
import json
import os
import re
from pathlib import Path

# 单个附件注入上文的文本上限（字符），超出截断，避免撑爆上下文
MAX_FILE_CHARS = 16_000
# 单张图片注入上限（字节）：base64 会膨胀 ~33%，且多数视觉 API 对大图有体积限制。
# 超出则放弃嵌入、降级为仅引用。可用 NA_MAX_IMAGE_BYTES 覆盖。
MAX_IMAGE_BYTES = int(os.environ.get("NA_MAX_IMAGE_BYTES", str(8 * 1024 * 1024)))
# 每轮最多嵌入的图片张数，超出部分仅引用。可用 NA_MAX_IMAGES_PER_TURN 覆盖。
MAX_IMAGES_PER_TURN = int(os.environ.get("NA_MAX_IMAGES_PER_TURN", "4"))

TEXT_EXTS = {".txt", ".md"}
HTML_EXTS = {".html"}
IMAGE_EXTS = {".png", ".jpg", ".jpeg"}


def _attachments_dir(conversations_dir: Path, conversation_id: str) -> Path:
    """对话附件目录，含路径穿越防护（与 routes/attachments.py 一致）。"""
    if "/" in conversation_id or "\\" in conversation_id or conversation_id in (".", ".."):
        raise ValueError(f"非法的 conversation_id：{conversation_id!r}")
    return conversations_dir / conversation_id / "attachments"


def _load_meta(attachments_dir: Path) -> list[dict]:
    """读 attachments.json 的附件列表；缺失/损坏返回空列表（不阻断）。"""
    meta_path = attachments_dir / "attachments.json"
    if not meta_path.exists():
        return []
    try:
        data = json.loads(meta_path.read_text(encoding="utf-8"))
        return list(data.get("attachments", []))
    except Exception:
        return []


def _decode_text(path: Path) -> str | None:
    """读取文本文件：UTF-8 → gb18030 → latin-1 依次回退，全部失败返回 None。"""
    for enc in ("utf-8", "gb18030", "latin-1"):
        try:
            return path.read_text(encoding=enc)
        except (UnicodeDecodeError, OSError):
            continue
    return None


def _strip_html(text: str) -> str:
    """正则级剥标签：去掉 script/style 内容块与其余标签，折叠空白。"""
    text = re.sub(r"<script[^>]*>.*?</script>", " ", text, flags=re.S | re.I)
    text = re.sub(r"<style[^>]*>.*?</style>", " ", text, flags=re.S | re.I)
    text = re.sub(r"<[^>]+>", " ", text)
    return re.sub(r"[ \t]+", " ", text).strip()


def _read_pdf(path: Path) -> str | None:
    try:
        from pypdf import PdfReader
    except Exception:
        return None
    try:
        reader = PdfReader(str(path))
        return "\n".join((page.extract_text() or "") for page in reader.pages).strip()
    except Exception:
        return None


def _read_docx(path: Path) -> str | None:
    try:
        from docx import Document
    except Exception:
        return None
    try:
        doc = Document(str(path))
        return "\n".join(p.text for p in doc.paragraphs).strip()
    except Exception:
        return None


def _read_xlsx(path: Path) -> str | None:
    try:
        import openpyxl
    except Exception:
        return None
    try:
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        lines: list[str] = []
        for ws in wb.worksheets:
            lines.append(f"[工作表 {ws.title}]")
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    lines.append(" | ".join(cells))
        return "\n".join(lines).strip()
    except Exception:
        return None


def read_attachment_text(entry: dict, attachments_dir: Path) -> str | None:
    """读取单个附件的可注入文本；不可读 / 解析失败返回 None（调用方降级为引用）。"""
    stored = Path(entry.get("stored_name") or "")
    target = attachments_dir / stored
    if not target.is_file():
        return None
    ext = Path(entry.get("name") or "").suffix.lower()
    if ext in TEXT_EXTS:
        return _decode_text(target)
    if ext in HTML_EXTS:
        raw = _decode_text(target)
        return _strip_html(raw) if raw is not None else None
    if ext == ".pdf":
        return _read_pdf(target)
    if ext == ".docx":
        return _read_docx(target)
    if ext == ".xlsx":
        return _read_xlsx(target)
    return None  # 图片与其他未支持格式：仅引用


def read_attachment_image_data_url(entry: dict, attachments_dir: Path) -> str | None:
    """读取图片附件为 base64 data URL（data:{mime};base64,{b64}）。

    文件缺失 / 体积超限 / 非图片扩展名返回 None（调用方降级为仅引用）。
    """
    ext = Path(entry.get("name") or "").suffix.lower()
    if ext not in IMAGE_EXTS:
        return None
    stored = Path(entry.get("stored_name") or "")
    target = attachments_dir / stored
    if not target.is_file():
        return None
    size = target.stat().st_size
    if size > MAX_IMAGE_BYTES:
        return None
    try:
        mime = {"png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg"}[ext.lstrip(".")]
        b64 = base64.b64encode(target.read_bytes()).decode("ascii")
    except OSError:
        return None
    return f"data:{mime};base64,{b64}"


def build_image_parts(
    conversations_dir: Path, conversation_id: str, attachment_ids: list[str] | None
) -> tuple[list[dict], set[str]]:
    """组装本轮图片输入（OpenAI 兼容 image_url content-parts）。

    返回 (parts, embedded_image_ids)：parts 为成功读为 data URL 的图片 content-parts；
    embedded_image_ids 为对应附件 id 集合（供文本块标注「图像已随消息输入」）。
    读取失败 / 超限 / 超过每轮张数上限的图片不在此列（由文本块降级为引用）。
    空 / 无图返回 ([], set())。
    """
    if not attachment_ids:
        return [], set()
    try:
        attachments_dir = _attachments_dir(conversations_dir, conversation_id)
    except ValueError:
        return [], set()

    id_set = set(attachment_ids)
    entries = [e for e in _load_meta(attachments_dir) if e.get("id") in id_set and e.get("name")]
    parts: list[dict] = []
    embedded_image_ids: set[str] = set()
    for e in entries:
        if len(parts) >= MAX_IMAGES_PER_TURN:
            break
        data_url = read_attachment_image_data_url(e, attachments_dir)
        if data_url:
            parts.append({"type": "image_url", "image_url": {"url": data_url}})
            embedded_image_ids.add(str(e.get("id")))
    return parts, embedded_image_ids


def build_attachment_context(
    conversations_dir: Path,
    conversation_id: str,
    attachment_ids: list[str] | None,
    *,
    embedded_image_ids: set[str] | None = None,
) -> str:
    """组装「附带文件」上下文块；空附件 / 无有效匹配返回空串（不注入）。

    embedded_image_ids：已作为图像输入（image_url part）随消息传给视觉模型的
    附件 id 集合；命中项在文本块中标注「图像已随消息输入」而非「仅引用」，
    避免模型看到重复的引用说明产生混淆。
    """
    if not attachment_ids:
        return ""
    try:
        attachments_dir = _attachments_dir(conversations_dir, conversation_id)
    except ValueError:
        return ""

    embedded = embedded_image_ids or set()
    id_set = set(attachment_ids)
    entries = [e for e in _load_meta(attachments_dir) if e.get("id") in id_set]
    if not entries:
        return ""

    lines = ["[本次对话附带文件]"]
    for e in entries:
        name = e.get("name", "") or "（未命名文件）"
        size_kb = (e.get("size") or 0) / 1024
        if e.get("id") in embedded:
            lines.append(f"· {name}（{size_kb:.1f}KB，图像已随消息输入）")
            continue
        text = read_attachment_text(e, attachments_dir)
        if text:
            if len(text) > MAX_FILE_CHARS:
                text = text[:MAX_FILE_CHARS] + "\n[内容过长已截断]"
            lines.append(f"· {name}：")
            lines.append(text)
        else:
            lines.append(f"· {name}（{size_kb:.1f}KB，无法解析内容，仅引用文件名）")
    return "\n".join(lines)
