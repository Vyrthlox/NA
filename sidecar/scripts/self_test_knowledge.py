#!/usr/bin/env python3
"""RAG 知识库自测：分词 / 切分 / BM25 / RRF 去重 / 级联重排 / 服务集成（自建 KB 全流程）。

用法（仓库根目录执行）：
    python -m sidecar.scripts.self_test_knowledge             # 检索真实、生成走 FakeLLM
    python -m sidecar.scripts.self_test_knowledge --real-llm  # 生成也走真实 API（需 temp/.env key）
    python -m sidecar.scripts.self_test_knowledge --heavy     # 额外加载高精度重排模型（默认关闭，显式开启；CPU 加载较慢）

高精度重排模型（bge-reranker-v2-m3）体积大且 CPU 上逐查询不可行，默认 SKIP（KN_RERANK_HEAVY=0）；
轻量重排（MiniLM）默认参与。

服务集成按 per-KB 模型跑：建库 → 上传（含子目录相对路径）→ kb 限定/跨库检索 → 生成 →
增量 reindex → 删文件 → 删库，全部断言「变更+重建一体」。
"""

from __future__ import annotations

import argparse
import shutil
import sys
import tempfile
from pathlib import Path

# 仓库根目录 = sidecar/scripts/self_test_knowledge.py 向上 3 级
REPO_ROOT = Path(__file__).resolve().parents[2]

for stream in (sys.stdout, sys.stderr):
    try:
        stream.reconfigure(encoding="utf-8")
    except (AttributeError, ValueError):
        pass

sys.path.insert(0, str(REPO_ROOT))

# 模型缓存落仓库 data/memory/models（跨模块复用，bge-small-zh 已缓存）
import os  # noqa: E402
os.environ.setdefault("HF_HOME", str(REPO_ROOT / "data" / "memory" / "models"))
os.environ.setdefault("HF_ENDPOINT", "https://hf-mirror.com")

import numpy as np  # noqa: E402

from sidecar.services.context.config import SidecarConfig  # noqa: E402
from sidecar.services.context.llm import FakeLLMClient  # noqa: E402
from sidecar.services.knowledge.bm25 import BM25Index  # noqa: E402
from sidecar.services.knowledge.chunker import chunk_document  # noqa: E402
from sidecar.services.knowledge.config import KnowledgeConfig  # noqa: E402
from sidecar.services.knowledge.rerank import CascadeReranker  # noqa: E402
from sidecar.services.knowledge.rrf import dedup_chunk_id, dedup_cosine, rrf_fuse  # noqa: E402
from sidecar.services.knowledge.service import KnowledgeService  # noqa: E402
from sidecar.services.knowledge.tokenizer import (  # noqa: E402
    extract_quoted, has_cjk, looks_like_identifier, tokenize,
)
from sidecar.services.knowledge.types import Chunk  # noqa: E402

FAIL = 0


def check(name: str, cond: bool, detail: object = "") -> None:
    global FAIL
    status = "PASS" if cond else "FAIL"
    if not cond:
        FAIL += 1
    text = str(detail) if detail else ""
    print(f"  [{status}] {name}{(' — ' + text) if text else ''}")


def _sample_note1() -> str:
    parts = ["<h1>光锥笔记产品需求</h1>"]
    for i in range(30):
        parts.append(
            f"<p>光锥笔记是一款本地优先的智能笔记应用，支持富文本编辑、AI 对话与知识库检索。"
            f"第{i}段：RAG 混合检索结合 BM25 关键词召回与 FAISS 向量召回，经 RRF 排名融合，"
            f"再用重排模型精排，最终把 Top 结果送入大模型生成回答。</p>"
        )
    parts.append("<h2>知识库技术细节</h2>")
    for i in range(10):
        parts.append(
            f"<p>索引使用 bge-small-zh 把文本编码为 512 维向量，chunk 控制在 256 到 512 个 token，"
            f"相邻块带 50 token 重叠，语义断裂点强制切分。第{i}条：向量检索用 FAISS，关键词检索用 BM25。</p>"
        )
    return "".join(parts)


def _sample_note2() -> str:
    parts = ["<h1>季度会议记录</h1>"]
    for i in range(12):
        parts.append(
            f"<p>本周例会讨论了进度：前端编辑器完成自动保存，sidecar 接入 AI 对话。"
            f"第{i}条纪要：预算 10 万，下周复盘，负责人小明。</p>"
        )
    return "".join(parts)


def main() -> None:
    parser = argparse.ArgumentParser(description="NA RAG 知识库自测")
    parser.add_argument("--real-llm", action="store_true", help="生成用真实 API（默认 FakeLLM）")
    parser.add_argument("--heavy", action="store_true", help="加载高精度重排模型（默认关闭；CPU 上逐查询加载较慢）")
    parser.add_argument("--data-dir", default=None, help="数据根目录（默认临时目录，结束清理）")
    parser.add_argument("--keep", action="store_true", help="保留临时数据目录")
    args = parser.parse_args()

    # --heavy 显式开启高精度重排（须在 KnowledgeConfig 构造前生效）
    if args.heavy:
        os.environ["KN_RERANK_HEAVY"] = "1"

    tmp = Path(args.data_dir) if args.data_dir else Path(tempfile.mkdtemp(prefix="na-knowledge-test-"))
    print("=" * 72)
    print("NA RAG 知识库自测")
    print(f"  数据目录: {tmp}")
    cfg = KnowledgeConfig(data_dir=tmp)
    print(f"  embed: {cfg.embed_model} ({cfg.embed_dims}维) | light: {cfg.rerank_light_model}")
    print(f"  heavy: {cfg.rerank_heavy_model} | BM25 短/长/精确 = {cfg.bm25_short}/{cfg.bm25_long}/{cfg.bm25_exact}")
    print("=" * 72)

    # ---------- 1. 分词 ----------
    print("\n[1] 分词器（tokenizer）")
    toks = tokenize("光锥笔记使用FAISS和BM25做混合检索，版本v1.2.3")
    check("中英混合分词", "光锥" in toks and "faiss" in toks and "bm25" in toks and "v1.2.3" in toks, toks[:12])
    check("小写归一", all(t == t.lower() for t in toks))
    check("has_cjk 识别中文", has_cjk("你好世界") and not has_cjk("hello world 123"))
    check("identifier 判定", looks_like_identifier("NA2024") and looks_like_identifier("BUG-123")
          and not looks_like_identifier("这是中文"))
    check("引号短语提取", extract_quoted('请介绍「RAG」') == "RAG" and extract_quoted("没有引号") is None)
    check("空文本分词", tokenize("") == [])

    # ---------- 2. 切分 ----------
    print("\n[2] 切分器（chunker，纯结构 + 语义校验占位）")
    doc_id = "方案/光锥笔记.html"
    chunks = chunk_document(doc_id, _sample_note1(), None, cfg)
    check("切出多个 chunk", len(chunks) >= 3, f"{len(chunks)} 个")
    check("token 不超上限", all(0 < c.tokens <= cfg.chunk_max_tokens for c in chunks),
          f"max={max(c.tokens for c in chunks)}")
    check("chunk_id 稳定前缀", all(c.chunk_id.startswith(doc_id + "::") for c in chunks), chunks[0].chunk_id)
    tech = [c for c in chunks if "技术细节" in c.text]
    check("h2 硬切（技术细节不混入前文段落）", bool(tech) and not any("第29段" in c.text for c in tech),
          f"{len(tech)} 个含技术细节")
    check("section 记录标题路径", any("技术细节" in (c.section or "") for c in tech),
          tech[0].section if tech else "")

    # ---------- 3. BM25 ----------
    print("\n[3] BM25（per-doc 参数 + exact 模式）")
    bm = BM25Index(cfg)
    ch_a = Chunk(chunk_id="a::0", doc_id="a", text="苹果 苹果 苹果 苹果 苹果", tokens=0)
    ch_b = Chunk(chunk_id="b::0", doc_id="b", text="苹果 香蕉 梨子", tokens=0)
    ch_a.tokens = len(tokenize(ch_a.text))
    ch_b.tokens = len(tokenize(ch_b.text))
    bm.build([ch_a, ch_b])
    auto = bm.search(tokenize("苹果"), top_k=2)
    exact = bm.search(tokenize("苹果"), top_k=2, exact=True)
    check("auto 模式词频加权（重复出现者靠前）", auto[0][0] == "a::0", str(auto))
    check("exact 模式 k1=0 平权", len(exact) == 2 and abs(exact[0][1] - exact[1][1]) < 1e-9, str(exact))
    check("empty query 返回空", bm.search([], top_k=2) == [])

    # ---------- 4. RRF + 去重 ----------
    print("\n[4] RRF 融合 + 去重")
    fused = rrf_fuse([[("c1", 1.0), ("c2", 0.5)], [("c1", 1.0), ("c3", 0.4)]], k=60)
    check("RRF 融合唯一性", {c for c, _ in fused} == {"c1", "c2", "c3"}, str(fused))
    ded = dedup_chunk_id([("c1", 0.5), ("c1", 0.8), ("c2", 0.4)])
    check("chunk_id 去重保最高分", ded == [("c1", 0.8), ("c2", 0.4)], str(ded))
    vecs = np.array([[1.0, 0.0], [0.0, 1.0], [0.99, 0.141]], dtype=np.float32)
    pos = {"c1": 0, "c2": 1, "c3": 2}
    sim = dedup_cosine([("c1", 0.5), ("c2", 0.4), ("c3", 0.3)], pos, vecs, 0.92)
    check("余弦>0.92 相似去重（c3 与 c1 相似被丢）", [c for c, _ in sim] == ["c1", "c2"], str(sim))

    # ---------- 5. 级联重排 ----------
    print("\n[5] 级联重排（轻量必测 / 高精度可选）")
    reranker = CascadeReranker(cfg)
    cands = [
        Chunk(chunk_id="r::0", doc_id="r", text="光锥笔记支持富文本编辑与 AI 对话。"),
        Chunk(chunk_id="r::1", doc_id="r", text="预算 10 万，下周复盘。"),
        Chunk(chunk_id="r::2", doc_id="r", text="RAG 混合检索用 BM25 与 FAISS。"),
        Chunk(chunk_id="r::3", doc_id="r", text="版本 v1.2.3 发布计划。"),
    ]
    if reranker.light_available():
        out = reranker.rerank("光锥笔记 富文本 AI 对话 知识库", cands, final_top=2)
        check("轻量重排后数量 ≤ final_top 且降序", 0 < len(out) <= 2 and all(
            out[i][1] >= out[i + 1][1] for i in range(len(out) - 1)), f"{len(out)} 个")
        check("重排后相关文档靠前", out[0][0].text.startswith("光锥笔记"), out[0][0].text[:24])
    else:
        print(f"  [SKIP] 轻量重排模型不可用：{reranker.light_error or '未知'}")
    if args.heavy:
        if reranker.heavy_available():
            out = reranker.rerank("光锥笔记 富文本 AI 对话", cands, final_top=3)
            check("高精度重排输出 ≤ final_top", 0 < len(out) <= 3)
        else:
            print(f"  [SKIP] 高精度重排模型不可用：{reranker.heavy_error or '未知'}")
    else:
        print("  [SKIP] 高精度重排默认跳过（--heavy 加载，首次下载 ~1GB）")

    # ---------- 6. 服务集成（自建 KB → 上传 → 检索 → 生成 → 删文件 → 删 KB） ----------
    print("\n[6] KnowledgeService 集成（per-KB：建库 → 上传 → 检索 → 生成 → 删除）")

    class _FakeRouter:
        def get(self, provider: str, depth: str):
            return FakeLLMClient()

    ctx_cfg = SidecarConfig(data_dir=tmp)
    svc = KnowledgeService(ctx_cfg, cfg, llm_router=_FakeRouter())
    svc._ready = True

    # 建库
    rec = svc.create_kb("测试知识库")
    kb_id = rec["id"]
    check("create_kb 返回 id/name", bool(kb_id) and rec["name"] == "测试知识库", str(rec))
    kbs0 = svc.list_kbs()
    check("list_kbs 含新库且 files=0", len(kbs0) == 1 and kbs0[0]["files"] == 0, str(kbs0))
    check("创建即建 files 目录", svc._ensure_kb(kb_id).cfg.files_dir.is_dir())

    # 上传（含子目录相对路径）→ 变更+重建一体
    uploads = [
        ("产品/光锥笔记.html", _sample_note1().encode("utf-8")),
        ("季度会议.html", _sample_note2().encode("utf-8")),
        ("老笔记.txt", "# 旧笔记\n\nRAG 知识库的架构设计与检索链路说明。".encode("utf-8")),
    ]
    stats = svc.upload_files(kb_id, uploads)
    check("upload 变更即重建（chunks>0）", stats["chunks"] >= 3, str(stats))
    check("upload written=3", stats["written"] == 3, str(stats))
    kb_cfg = svc._ensure_kb(kb_id).cfg
    check("每 KB 独立索引落盘", kb_cfg.chunks_path.exists() and kb_cfg.faiss_path.with_suffix(".npy").exists())
    files = svc.list_kb_files(kb_id)
    check("list_kb_files 保留子目录 rel", len(files) == 3
          and any(f["rel"] == "产品/光锥笔记.html" for f in files), [f["rel"] for f in files])
    kbs1 = svc.list_kbs()
    check("kb 摘要 files/chunks/indexed", kbs1[0]["files"] == 3 and kbs1[0]["chunks"] > 0 and kbs1[0]["indexed"],
          str(kbs1[0]))

    # kb 维度检索
    res = svc.search("光锥笔记有哪些功能", kb_id=kb_id)
    check("kb 限定 search 非空", len(res) >= 1, f"{len(res)} 条")
    check("search 结果字段完整", all(r.chunk_id and r.doc_id and r.text for r in res))
    check("search ≤ final_top_max", len(res) <= cfg.final_top_max, f"{len(res)}/{cfg.final_top}")
    res_exact = svc.search('请介绍「RAG」', kb_id=kb_id)
    check("引号短语触发 exact 检索", len(res_exact) >= 1 and any("exact" in r.sources for r in res_exact),
          [r.sources for r in res_exact])

    # 全局 search（跨库合并） vs kb 限定
    rec2 = svc.create_kb("第二库")
    svc.upload_files(rec2["id"], [("另文.html", _sample_note1().encode("utf-8"))])
    g = svc.search("光锥笔记")
    check("缺省 search 跨库合并", len(g) >= 1 and all(r.doc_id for r in g), f"{len(g)} 条")
    g1 = svc.search("光锥笔记", kb_id=rec2["id"])
    check("kb 限定只搜该库", len(g1) >= 1 and all(r.doc_id == "另文.html" for r in g1), [r.doc_id for r in g1])

    # 增量 reindex：无变更全跳过；直接改磁盘文件 → 仅该文档更新
    stats2 = svc.reindex_kb(kb_id, incremental=True)
    check("增量 reindex 无变更跳过", stats2["updated"] == 0 and stats2["chunks"] == stats["chunks"], str(stats2))
    (kb_cfg.files_dir / "季度会议.html").write_text(
        _sample_note2() + "<p>追加：新增预算 20 万。</p>", encoding="utf-8")
    stats3 = svc.reindex_kb(kb_id, incremental=True)
    check("增量 reindex 仅更新变更文档", stats3["updated"] == 1, str(stats3))

    # 生成（kb 限定）
    out = svc.query("光锥笔记有哪些功能？", kb_id=kb_id, provider="deepseek", depth="fast")
    check("query 返回 answer", "FakeLLM" in out["answer"], out["answer"][:40])
    check("query 返回 citations", len(out["citations"]) >= 1 and "doc_id" in out["citations"][0])

    if args.real_llm:
        from sidecar.services.context.llm import LLMRouter  # noqa: E402
        svc.router = LLMRouter(ctx_cfg)
        try:
            real = svc.query("光锥笔记的 RAG 检索是怎么做的？", kb_id=kb_id)
            check("真实 LLM 生成（引用标注）", len(real["answer"]) > 10, real["answer"][:60])
        except Exception as e:
            print(f"  [SKIP] 真实 LLM 不可用：{e}")

    # 删除文件 → 变更+重建（老笔记.txt 恰好 1 个 chunk）
    before_files = len(svc.list_kb_files(kb_id))
    d_stats = svc.delete_kb_file(kb_id, "老笔记.txt")
    after_files = len(svc.list_kb_files(kb_id))
    check("delete_kb_file 删文件并重建", before_files == after_files + 1
          and d_stats["chunks"] == stats3["chunks"] - 1,
          f"{before_files}→{after_files} chunks={stats3['chunks']}→{d_stats['chunks']}")
    try:
        svc.delete_kb_file(kb_id, "不存在.md")
        check("删除不存在文件抛 FileNotFoundError", False)
    except FileNotFoundError:
        check("删除不存在文件抛 FileNotFoundError", True)

    # 删除 KB → 目录 + 注册表
    kb_dir = kb_cfg.data_dir
    ok = svc.delete_kb(kb_id)
    check("delete_kb 移除注册表 + 目录", ok and not kb_dir.exists()
          and all(r["id"] != kb_id for r in svc.registry.list()))
    check("delete_kb 后 list_kbs 仅剩另一库", len(svc.list_kbs()) == 1, str([r["name"] for r in svc.list_kbs()]))

    # 路径穿越防护
    try:
        svc.upload_files(rec2["id"], [("../escape.html", b"x")])
        check("上传拒绝 ../ 穿越", False)
    except ValueError:
        check("上传拒绝 ../ 穿越", True)

    # ---------- 7. 格式白名单 + 多格式提取（txt/csv/json/docx/pdf） ----------
    print("\n[7] 上传格式白名单 + 多格式文本提取")

    def _make_pdf(text: str) -> bytes:
        """构造单页最小合法 PDF（Helvetica + 一行 ASCII 文本，xref 偏移程序化计算）。"""
        content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("ascii")
        objs = [
            b"<< /Type /Catalog /Pages 2 0 R >>",
            b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
            b"/Resources << /Font << /F1 5 0 R >> >> >>",
            b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        ]
        out = bytearray(b"%PDF-1.4\n")
        offsets = []
        for i, body in enumerate(objs, start=1):
            offsets.append(len(out))
            out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
        xref_pos = len(out)
        out += f"xref\n0 {len(objs) + 1}\n".encode()
        out += b"0000000000 65535 f \n"
        for off in offsets:
            out += f"{off:010d} 00000 n \n".encode()
        out += b"trailer\n" + f"<< /Size {len(objs) + 1} /Root 1 0 R >>\n".encode()
        out += f"startxref\n{xref_pos}\n%%EOF".encode()
        return bytes(out)

    fmt_kb = svc.create_kb("格式库")["id"]
    txt_b = "格式自检：青柠气泡水是本次测试的唯一锚点文本。".encode("utf-8")
    csv_b = "项目,编号\n灯塔计划,LM-7788\n灯塔计划,LM-7789\n".encode("utf-8")
    json_b = '{"项目": "极光档案", "编号": "JG-2026", "标签": ["测试", "格式"]}'.encode("utf-8")
    md_b = "# 需求文档\n\n## 目标\n\n青柠气泡水是本次测试的唯一锚点文本。\n\n- 要点一\n- 要点二\n".encode("utf-8")
    fmt_uploads = [
        ("fmt/样例.txt", txt_b),
        ("fmt/说明.md", md_b),
        ("fmt/清单.csv", csv_b),
        ("fmt/配置.json", json_b),
        ("fmt/页面.htm", "<h1>htm 样例</h1><p>信天翁航线说明。</p>".encode("utf-8")),
    ]
    # docx 夹具（python-docx 生成：Heading + 段落 + 表格）
    try:
        from docx import Document as _DocxDocument  # noqa: E402
        import io as _io
        _d = _DocxDocument()
        _d.add_heading("格式样例", level=1)
        _d.add_paragraph("白皮书核心结论：琥珀色信天翁是测试锚点。")
        _t = _d.add_table(rows=2, cols=2)
        _t.cell(0, 0).text = "项目"; _t.cell(0, 1).text = "值"
        _t.cell(1, 0).text = "灯塔计划"; _t.cell(1, 1).text = "LM-7788"
        _buf = _io.BytesIO(); _d.save(_buf)
        fmt_uploads.append(("fmt/报告.docx", _buf.getvalue()))
        has_docx = True
    except ImportError:
        print("  [SKIP] python-docx 未安装，跳过 docx 夹具")
        has_docx = False
    # pdf 夹具（手写最小 PDF；提取依赖 pypdf）
    pdf_text = "amber albatross fixture phrase 9911"
    fmt_uploads.append(("fmt/扫描.pdf", _make_pdf(pdf_text)))
    try:
        import pypdf  # noqa: F401, E402
        has_pypdf = True
    except ImportError:
        print("  [SKIP] pypdf 未安装，pdf 上传将计入 failed（不断言命中）")
        has_pypdf = False

    fmt_stats = svc.upload_files(fmt_kb, fmt_uploads)
    expect_chunks = fmt_stats["failed"] == 0
    check("多格式上传 written 全量落盘", fmt_stats["written"] == len(fmt_uploads), str(fmt_stats))
    if has_pypdf:
        check("多格式上传 failed=0（含 docx/pdf 提取）", expect_chunks, str(fmt_stats))
    check("格式库 chunks>0", fmt_stats["chunks"] > 0, str(fmt_stats))
    fmt_files = svc.list_kb_files(fmt_kb)
    check("list_kb_files 含全部格式", {f["rel"] for f in fmt_files}
          == {rel for rel, _ in fmt_uploads}, [f["rel"] for f in fmt_files])

    def _hit(query: str, rel_suffix: str) -> bool:
        return any(r.doc_id.endswith(rel_suffix) for r in svc.search(query, kb_id=fmt_kb))

    check("txt 内容可检索", _hit("青柠气泡水", ".txt"))
    check("md 纯文本可检索", _hit("青柠气泡水", ".md"))
    check("csv 行内容可检索", _hit("灯塔计划", ".csv"))
    check("json 值可检索", _hit("极光档案", ".json"))
    check("htm 走结构解析可检索", _hit("信天翁航线", ".htm"))
    if has_docx:
        check("docx 段落可检索", _hit("琥珀色信天翁", ".docx"))
    if has_pypdf:
        check("pdf 文本可检索", _hit("albatross", ".pdf"))
    stats_fmt2 = svc.reindex_kb(fmt_kb, incremental=True)
    check("格式库增量 reindex 无变更跳过", stats_fmt2["updated"] == 0, str(stats_fmt2))

    # 白名单拒绝：单个非法 / 混批（整批拒绝，不落盘任何文件）
    try:
        svc.upload_files(fmt_kb, [("病毒.exe", b"MZ")])
        check("上传拒绝白名单外扩展名", False)
    except ValueError as e:
        check("上传拒绝白名单外扩展名", "不支持的文件格式" in str(e), str(e)[:60])
    n_before = len(svc.list_kb_files(fmt_kb))
    try:
        svc.upload_files(fmt_kb, [("合法.txt", b"ok"), ("图片.png", b"\x89PNG")])
        check("混批含非法整体拒绝", False)
    except ValueError:
        check("混批含非法整体拒绝", len(svc.list_kb_files(fmt_kb)) == n_before, "未落盘")
    try:
        svc.upload_files(fmt_kb, [("无扩展名", b"x")])
        check("无扩展名拒绝", False)
    except ValueError:
        check("无扩展名拒绝", True)

    # 损坏 pdf：落盘但提取失败 → 计入 failed，服务不崩
    bad_stats = svc.upload_files(fmt_kb, [("fmt/损坏.pdf", b"this is not a pdf")])
    check("损坏 pdf 计入 failed 不崩", bad_stats["failed"] >= 1, str(bad_stats))
    check("损坏 pdf 不影响既有 chunks", bad_stats["chunks"] >= fmt_stats["chunks"], str(bad_stats))

    svc.delete_kb(fmt_kb)
    svc.shutdown()

    # ---------- 汇总 ----------
    print("\n" + "=" * 72)
    if FAIL:
        print(f"结果：{FAIL} 项 FAIL（其余通过）")
        sys.exit(1)
    print("结果：全部 PASS")
    if not args.keep and not args.data_dir:
        shutil.rmtree(tmp, ignore_errors=True)
        print(f"（临时目录已清理：{tmp}）")
    elif args.data_dir:
        print(f"（数据保留于：{tmp}）")


if __name__ == "__main__":
    main()
